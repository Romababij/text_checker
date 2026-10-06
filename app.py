"""AutoDoc multilingual Excel content validator — Streamlit UI."""

from __future__ import annotations

import io
from typing import Optional

import pandas as pd
import streamlit as st
from openpyxl import load_workbook
from openpyxl.styles import PatternFill

from validator import (
    ContentValidator,
    RowValidationResult,
    build_full_export_table,
    build_results_table,
)

st.set_page_config(
    page_title="AutoDoc Content Validator",
    page_icon="🛡️",
    layout="wide",
)

PROBLEM_FILL = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")


@st.cache_resource
def get_validator() -> ContentValidator:
    return ContentValidator()


def export_validated_excel(
    df: pd.DataFrame,
    results: list[RowValidationResult],
    text_column: str,
) -> bytes:
    out_df = df.copy()
    by_row = {r.row_index: r for r in results}
    statuses, l1s, l2s, l3s = [], [], [], []
    for idx in df.index:
        excel_row = int(idx) + 2
        r = by_row.get(excel_row)
        if r is None:
            statuses.append("OK")
            l1s.append("")
            l2s.append("")
            l3s.append("")
        else:
            statuses.append(r.status_summary)
            l1s.append(r.l1_detail)
            l2s.append(r.l2_detail)
            l3s.append(r.l3_detail)
    out_df["Validation_Status"] = statuses
    out_df["L1_Script"] = l1s
    out_df["L2_Spelling"] = l2s
    out_df["L3_US_GB"] = l3s

    buffer = io.BytesIO()
    out_df.to_excel(buffer, index=False, engine="openpyxl")
    buffer.seek(0)

    wb = load_workbook(buffer)
    ws = wb.active
    headers = [cell.value for cell in ws[1]]
    try:
        text_col_idx = headers.index(text_column) + 1
    except ValueError:
        text_col_idx = None

    for r in results:
        if r.is_clean or text_col_idx is None:
            continue
        ws.cell(row=r.row_index, column=text_col_idx).fill = PROBLEM_FILL

    out_buffer = io.BytesIO()
    wb.save(out_buffer)
    return out_buffer.getvalue()


def main() -> None:
    st.title("AutoDoc — Multilingual Content Validator")
    st.caption("L1 script/homoglyphs · L2 spelling & Lingua · L3 US vs British English")

    uploaded = st.file_uploader("Excel file (.xlsx)", type=["xlsx"])
    if uploaded is None:
        st.info("Upload a `.xlsx` file to begin.")
        return

    try:
        df = pd.read_excel(uploaded, engine="openpyxl")
    except Exception as exc:
        st.error(f"Could not read Excel: {exc}")
        return

    if df.empty:
        st.warning("No data rows in file.")
        return

    columns = list(df.columns.astype(str))
    c1, c2 = st.columns(2)
    with c1:
        lang_column = st.selectbox("Language column", columns, index=0)
    with c2:
        text_column = st.selectbox(
            "Text column", columns, index=min(1, len(columns) - 1)
        )

    if st.button("Run validation", type="primary"):
        with st.spinner("Validating…"):
            results = get_validator().validate_dataframe(df, lang_column, text_column)
        st.session_state["results"] = results
        st.session_state["source_df"] = df
        st.session_state["lang_column"] = lang_column
        st.session_state["text_column"] = text_column

    results: Optional[list[RowValidationResult]] = st.session_state.get("results")
    if not results:
        st.caption(f"{len(df)} rows loaded. Run validation to see results.")
        return

    total = len(results)
    clean = sum(1 for r in results if r.is_clean)
    st.metric("Total rows", total)
    m2, m3 = st.columns(2)
    m2.metric("Clean", clean)
    m3.metric("Issues", total - clean)

    problem_df = build_results_table(results)
    langs = sorted({r.lang_code_raw for r in results if not r.is_clean})
    f_lang = st.multiselect("Filter language", langs, default=[])
    f_l1 = st.checkbox("Only rows with L1 issues", value=False)
    f_l2 = st.checkbox("Only rows with L2 issues", value=False)
    f_l3 = st.checkbox("Only rows with L3 issues", value=False)

    filtered = problem_df
    if not filtered.empty:
        if f_lang:
            filtered = filtered[filtered["Language"].isin(f_lang)]
        if f_l1:
            filtered = filtered[filtered["L1 (Script/Homoglyph)"].astype(str).str.len() > 0]
        if f_l2:
            filtered = filtered[filtered["L2 (Spelling/Lingua)"].astype(str).str.len() > 0]
        if f_l3:
            filtered = filtered[filtered["L3 (US vs GB)"].astype(str).str.len() > 0]

    st.subheader("Problem rows")
    if filtered.empty:
        st.success("No issues (or none match filters).")
    else:
        st.dataframe(filtered, use_container_width=True, hide_index=True)

    st.subheader("Export")
    src = st.session_state.get("source_df", df)
    text_col = st.session_state.get("text_column", text_column)
    full_report = build_full_export_table(results)

    try:
        xlsx = export_validated_excel(src, results, text_col)
        st.download_button(
            "Download validated_articles.xlsx",
            data=xlsx,
            file_name="validated_articles.xlsx",
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
    except Exception as exc:
        st.error(f"Excel export failed: {exc}")

    csv_buf = full_report.to_csv(index=False).encode("utf-8")
    st.download_button(
        "Download validation_report.csv",
        data=csv_buf,
        file_name="validation_report.csv",
        mime="text/csv",
    )


if __name__ == "__main__":
    main()
